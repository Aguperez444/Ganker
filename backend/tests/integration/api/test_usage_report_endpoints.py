import csv
import io
from datetime import datetime

import pytest
from openpyxl import load_workbook

from app.infrastructure.database.models.character_orm import CharacterORM
from app.infrastructure.database.models.character_priority_orm import CharacterPriorityORM
from app.infrastructure.database.models.conversation_orm import ConversationORM
from app.infrastructure.database.models.game_profile_orm import GameProfileORM
from app.infrastructure.database.models.rank_orm import RankORM
from app.infrastructure.database.models.role_orm import RoleORM
from app.infrastructure.database.models.role_profile_orm import RoleProfileORM
from app.infrastructure.database.models.team_orm import TeamORM
from app.infrastructure.database.models.user_orm import UserORM
from app.infrastructure.database.models.videogame_orm import VideogameORM


@pytest.fixture
def usage_data(test_db_session, password_hasher):
    """Crea un videojuego con rangos, personajes, perfiles de jugador y equipos para probar el reporte de uso."""
    lol = VideogameORM(name="League of Legends", icon_url="/media/lol.png", rank_per_role=True)
    valorant = VideogameORM(name="Valorant", icon_url="/media/val.png", rank_per_role=False)
    test_db_session.add_all([lol, valorant])
    test_db_session.flush()

    gold = RankORM(name="Gold", value=1000, videogame_id=lol.videogame_id, icon_url="/gold.png")
    platinum = RankORM(name="Platinum", value=2000, videogame_id=lol.videogame_id, icon_url="/plat.png")
    diamond = RankORM(name="Diamond", value=3000, videogame_id=lol.videogame_id, icon_url="/dia.png")
    test_db_session.add_all([gold, platinum, diamond])

    mid = RoleORM(name="Mid", videogame_id=lol.videogame_id, icon_url="/mid.png")
    test_db_session.add(mid)

    ahri = CharacterORM(name="Ahri", videogame_id=lol.videogame_id, icon_url="/ahri.png")
    yasuo = CharacterORM(name="Yasuo", videogame_id=lol.videogame_id, icon_url="/yasuo.png")
    test_db_session.add_all([ahri, yasuo])
    test_db_session.flush()

    players = [
        UserORM(
            name=f"Player {index}",
            username=f"player{index}",
            mail=f"player{index}@example.com",
            password_hash=password_hasher.hash_password("Password123"),
            role="player",
        )
        for index in range(3)
    ]
    test_db_session.add_all(players)
    test_db_session.flush()

    now = datetime(2026, 6, 15, 12, 0, 0)
    profiles = []
    for player in players:
        profile = GameProfileORM(player_id=player.user_id, videogame_id=lol.videogame_id, created_at=now)
        test_db_session.add(profile)
        profiles.append(profile)
    test_db_session.flush()

    # Dos jugadores en Gold, uno en Platinum
    test_db_session.add_all([
        RoleProfileORM(game_profile_id=profiles[0].game_profile_id, role_id=mid.role_id, rank_id=gold.rank_id),
        RoleProfileORM(game_profile_id=profiles[1].game_profile_id, role_id=mid.role_id, rank_id=gold.rank_id),
        RoleProfileORM(game_profile_id=profiles[2].game_profile_id, role_id=mid.role_id, rank_id=platinum.rank_id),
    ])
    # Dos jugadores con Ahri, uno con Yasuo
    test_db_session.add_all([
        CharacterPriorityORM(game_profile_id=profiles[0].game_profile_id, character_id=ahri.character_id, priority=1),
        CharacterPriorityORM(game_profile_id=profiles[1].game_profile_id, character_id=ahri.character_id, priority=1),
        CharacterPriorityORM(game_profile_id=profiles[2].game_profile_id, character_id=yasuo.character_id, priority=1),
    ])

    # Dos equipos con rango mínimo Gold (cada uno en su propia conversación)
    for name in ("Team Alpha", "Team Beta"):
        conversation = ConversationORM()
        test_db_session.add(conversation)
        test_db_session.flush()
        test_db_session.add(
            TeamORM(
                name=name,
                allow_other_regions=False,
                videogame_id=lol.videogame_id,
                min_rank_id=gold.rank_id,
                max_rank_id=diamond.rank_id,
                conversation_id=conversation.conversation_id,
                created_at=now,
            )
        )

    test_db_session.commit()
    return {
        "lol": lol,
        "valorant": valorant,
        "gold": gold,
        "platinum": platinum,
        "diamond": diamond,
        "ahri": ahri,
        "yasuo": yasuo,
        "players": players,
        "profiles": profiles,
    }


class TestGameUsageReportEndpointsIntegration:

    def _generate(self, client, headers, videogame_id, analysis_type="ranks", **extra):
        payload = {"videogame_id": videogame_id, "analysis_type": analysis_type, **extra}
        return client.post("/api/v1/reports/game-usage", json=payload, headers=headers)

    # ---------------------------------------------------------
    # POST /api/v1/reports/game-usage
    # ---------------------------------------------------------

    def test_generate_ranks_report_sorted_with_percentages(self, client, admin_auth_headers, usage_data):
        response = self._generate(client, admin_auth_headers, usage_data["lol"].videogame_id, "ranks")

        assert response.status_code == 201
        data = response.json()
        assert data["report_id"] is not None
        assert data["analysis_type"] == "ranks"
        assert data["videogame_name"] == "League of Legends"
        assert data["total"] == 3
        assert data["has_previous_report"] is False
        assert data["previous_report_id"] is None

        assert [item["name"] for item in data["items"]] == ["Gold", "Platinum"]
        gold = data["items"][0]
        assert gold["count"] == 2
        assert gold["percentage"] == 66.67
        assert gold["variation"] is None
        assert gold["trend"] is None

    def test_generate_characters_report(self, client, admin_auth_headers, usage_data):
        response = self._generate(client, admin_auth_headers, usage_data["lol"].videogame_id, "characters")

        assert response.status_code == 201
        data = response.json()
        assert data["analysis_type"] == "characters"
        assert [item["name"] for item in data["items"]] == ["Ahri", "Yasuo"]
        assert data["items"][0]["count"] == 2

    def test_generate_teams_report_with_rank_bands(self, client, admin_auth_headers, usage_data):
        response = self._generate(
            client,
            admin_auth_headers,
            usage_data["lol"].videogame_id,
            "teams",
            rank_bands=[{
                "min_rank_id": usage_data["gold"].rank_id,
                "max_rank_id": usage_data["diamond"].rank_id,
            }],
        )

        assert response.status_code == 201
        data = response.json()
        assert data["analysis_type"] == "teams"
        assert data["total"] == 2
        assert len(data["items"]) == 1
        band = data["items"][0]
        assert band["name"] == "Rango Gold - Diamond"
        assert band["rank_range"] == "1000 - 3000"
        assert band["category_id"] is None
        assert band["count"] == 2
        assert band["percentage"] == 100.0

    def test_generate_teams_without_rank_bands_fails(self, client, admin_auth_headers, usage_data):
        response = self._generate(client, admin_auth_headers, usage_data["lol"].videogame_id, "teams")

        assert response.status_code == 400
        assert "franjas de rangos" in response.json().get("error", "").lower()

    def test_generate_without_data_notifies_no_usage_data(self, client, admin_auth_headers, usage_data):
        response = self._generate(client, admin_auth_headers, usage_data["valorant"].videogame_id, "ranks")

        assert response.status_code == 404
        assert "no hay datos" in response.json().get("error", "").lower()

    def test_generate_invalid_analysis_type_fails(self, client, admin_auth_headers, usage_data):
        response = self._generate(client, admin_auth_headers, usage_data["lol"].videogame_id, "pizza")

        assert response.status_code == 400
        assert "no es válido" in response.json().get("error", "").lower()

    def test_generate_videogame_not_found(self, client, admin_auth_headers, usage_data):
        response = self._generate(client, admin_auth_headers, 999999, "ranks")

        assert response.status_code == 404

    def test_generate_forbidden_for_player(self, client, player_auth_headers, usage_data):
        response = self._generate(client, player_auth_headers, usage_data["lol"].videogame_id, "ranks")
        assert response.status_code == 403

    def test_generate_unauthorized_without_token(self, client, usage_data):
        response = self._generate(client, {}, usage_data["lol"].videogame_id, "ranks")
        assert response.status_code == 401

    def test_generate_with_previous_report_returns_variation(self, client, admin_auth_headers, usage_data):
        first = self._generate(client, admin_auth_headers, usage_data["lol"].videogame_id, "ranks")
        assert first.status_code == 201

        second = self._generate(client, admin_auth_headers, usage_data["lol"].videogame_id, "ranks")
        assert second.status_code == 201
        data = second.json()

        assert data["has_previous_report"] is True
        assert data["previous_report_id"] == first.json()["report_id"]
        assert all(item["variation"] == 0 for item in data["items"])
        assert all(item["trend"] == "estable" for item in data["items"])

    def test_generate_start_before_previous_end_fails(self, client, admin_auth_headers, usage_data):
        first = self._generate(
            client, admin_auth_headers, usage_data["lol"].videogame_id, "ranks",
            start_date="2026-06-01", end_date="2026-06-30",
        )
        assert first.status_code == 201

        second = self._generate(
            client, admin_auth_headers, usage_data["lol"].videogame_id, "ranks",
            start_date="2026-06-15", end_date="2026-06-20",
        )
        assert second.status_code == 400
        assert "fecha inicial" in second.json().get("error", "").lower()

    def test_generate_characters_with_date_filter(self, client, admin_auth_headers, usage_data, test_db_session):
        old_player = UserORM(
            name="Old Player", username="oldplayer", mail="old@example.com",
            password_hash="irrelevant", role="player",
        )
        test_db_session.add(old_player)
        test_db_session.flush()
        old_profile = GameProfileORM(
            player_id=old_player.user_id,
            videogame_id=usage_data["lol"].videogame_id,
            created_at=datetime(2026, 1, 1, 12, 0, 0),
        )
        test_db_session.add(old_profile)
        test_db_session.flush()
        test_db_session.add(
            CharacterPriorityORM(
                game_profile_id=old_profile.game_profile_id,
                character_id=usage_data["yasuo"].character_id,
                priority=1,
            )
        )
        test_db_session.commit()

        response = self._generate(
            client, admin_auth_headers, usage_data["lol"].videogame_id, "characters",
            start_date="2026-06-01", end_date="2026-06-30",
        )

        assert response.status_code == 201
        data = response.json()
        yasuo = next(item for item in data["items"] if item["name"] == "Yasuo")
        # El perfil de enero queda fuera del rango, así que Yasuo cuenta solo 1
        assert yasuo["count"] == 1

    # ---------------------------------------------------------
    # GET /api/v1/reports/game-usage
    # ---------------------------------------------------------

    def test_list_reports_after_generation(self, client, admin_auth_headers, usage_data):
        first = self._generate(client, admin_auth_headers, usage_data["lol"].videogame_id, "ranks").json()
        second = self._generate(client, admin_auth_headers, usage_data["lol"].videogame_id, "characters").json()

        response = client.get("/api/v1/reports/game-usage", headers=admin_auth_headers)

        assert response.status_code == 200
        reports = response.json()["reports"]
        assert [report["report_id"] for report in reports] == [second["report_id"], first["report_id"]]
        assert reports[0]["analysis_type"] == "characters"
        assert reports[0]["videogame_name"] == "League of Legends"

    def test_list_reports_filters_by_type(self, client, admin_auth_headers, usage_data):
        videogame_id = usage_data["lol"].videogame_id
        self._generate(client, admin_auth_headers, videogame_id, "ranks")
        self._generate(client, admin_auth_headers, videogame_id, "characters")

        ranks = client.get(
            f"/api/v1/reports/game-usage?videogame_id={videogame_id}&analysis_type=ranks",
            headers=admin_auth_headers,
        )

        assert ranks.status_code == 200
        reports = ranks.json()["reports"]
        assert len(reports) == 1
        assert reports[0]["analysis_type"] == "ranks"

    def test_list_reports_empty(self, client, admin_auth_headers, usage_data):
        response = client.get("/api/v1/reports/game-usage", headers=admin_auth_headers)

        assert response.status_code == 200
        assert response.json()["reports"] == []

    def test_list_reports_invalid_analysis_type(self, client, admin_auth_headers, usage_data):
        response = client.get("/api/v1/reports/game-usage?analysis_type=bogus", headers=admin_auth_headers)
        assert response.status_code == 400

    def test_list_reports_forbidden_for_player(self, client, player_auth_headers, usage_data):
        response = client.get("/api/v1/reports/game-usage", headers=player_auth_headers)
        assert response.status_code == 403

    # ---------------------------------------------------------
    # GET /api/v1/reports/game-usage/analysis-types
    # ---------------------------------------------------------

    def test_analysis_types_catalog(self, client, admin_auth_headers, usage_data):
        response = client.get("/api/v1/reports/game-usage/analysis-types", headers=admin_auth_headers)

        assert response.status_code == 200
        analysis_types = response.json()["analysis_types"]
        assert [entry["value"] for entry in analysis_types] == ["ranks", "characters", "teams"]
        by_value = {entry["value"]: entry for entry in analysis_types}
        assert by_value["teams"]["requires_rank_bands"] is True
        assert by_value["ranks"]["requires_rank_bands"] is False
        assert by_value["characters"]["requires_rank_bands"] is False

    def test_analysis_types_forbidden_for_player(self, client, player_auth_headers, usage_data):
        response = client.get("/api/v1/reports/game-usage/analysis-types", headers=player_auth_headers)
        assert response.status_code == 403

    # ---------------------------------------------------------
    # GET /api/v1/reports/game-usage/last
    # ---------------------------------------------------------

    def test_last_report_without_reports(self, client, admin_auth_headers, usage_data):
        response = client.get(
            f"/api/v1/reports/game-usage/last?analysis_type=ranks&videogame_id={usage_data['valorant'].videogame_id}",
            headers=admin_auth_headers,
        )

        assert response.status_code == 200
        data = response.json()
        assert data["has_report"] is False
        assert data["report_id"] is None
        assert data["suggested_start_date"] is None

    def test_last_report_after_generation(self, client, admin_auth_headers, usage_data):
        videogame_id = usage_data["lol"].videogame_id
        generated = self._generate(
            client, admin_auth_headers, videogame_id, "ranks",
            start_date="2026-06-01", end_date="2026-06-30",
        ).json()

        response = client.get(
            f"/api/v1/reports/game-usage/last?analysis_type=ranks&videogame_id={videogame_id}",
            headers=admin_auth_headers,
        )

        assert response.status_code == 200
        data = response.json()
        assert data["has_report"] is True
        assert data["report_id"] == generated["report_id"]
        assert data["end_date"].startswith("2026-06-30")
        assert data["suggested_start_date"] == "2026-07-01"

    def test_last_report_invalid_analysis_type(self, client, admin_auth_headers, usage_data):
        response = client.get(
            f"/api/v1/reports/game-usage/last?analysis_type=nope&videogame_id={usage_data['lol'].videogame_id}",
            headers=admin_auth_headers,
        )
        assert response.status_code == 400

    # ---------------------------------------------------------
    # GET /api/v1/reports/game-usage/{report_id}/export
    # ---------------------------------------------------------

    def test_export_report_csv_integrity(self, client, admin_auth_headers, usage_data):
        report_id = self._generate(client, admin_auth_headers, usage_data["lol"].videogame_id, "ranks").json()["report_id"]

        exported = client.get(f"/api/v1/reports/game-usage/{report_id}/export", headers=admin_auth_headers)

        assert exported.status_code == 200
        assert "text/csv" in exported.headers["content-type"]
        assert "attachment" in exported.headers["content-disposition"]

        rows = list(csv.reader(io.StringIO(exported.text)))
        assert rows[0] == ["RANGO", "JUGADORES", "PORCENTAJE", "VARIACION", "VARIACION %", "TENDENCIA"]
        data_rows = {
            row[0]: int(row[1]) for row in rows[1:] if len(row) == 6 and row[0] not in ("", "TOTAL")
        }
        assert data_rows["Gold"] == 2
        assert data_rows["Platinum"] == 1
        total_row = next(row for row in rows if row and row[0] == "TOTAL")
        assert int(total_row[1]) == 3

    def test_export_report_pdf(self, client, admin_auth_headers, usage_data):
        report_id = self._generate(client, admin_auth_headers, usage_data["lol"].videogame_id, "ranks").json()["report_id"]

        exported = client.get(
            f"/api/v1/reports/game-usage/{report_id}/export?report_output_format=pdf",
            headers=admin_auth_headers,
        )

        assert exported.status_code == 200
        assert exported.headers["content-type"] == "application/pdf"
        assert exported.headers["content-disposition"].endswith(".pdf")
        assert exported.content.startswith(b"%PDF")

    def test_export_report_excel(self, client, admin_auth_headers, usage_data):
        report_id = self._generate(client, admin_auth_headers, usage_data["lol"].videogame_id, "ranks").json()["report_id"]

        exported = client.get(
            f"/api/v1/reports/game-usage/{report_id}/export?report_output_format=excel",
            headers=admin_auth_headers,
        )

        assert exported.status_code == 200
        assert "spreadsheetml" in exported.headers["content-type"]
        assert exported.content.startswith(b"PK")
        worksheet = load_workbook(io.BytesIO(exported.content)).active
        values = [cell.value for row in worksheet.iter_rows() for cell in row]
        assert "Gold" in values
        assert "Platinum" in values

    def test_export_report_not_found(self, client, admin_auth_headers):
        response = client.get("/api/v1/reports/game-usage/99999/export", headers=admin_auth_headers)
        assert response.status_code == 404

    def test_export_report_unsupported_format(self, client, admin_auth_headers, usage_data):
        report_id = self._generate(client, admin_auth_headers, usage_data["lol"].videogame_id, "ranks").json()["report_id"]

        response = client.get(
            f"/api/v1/reports/game-usage/{report_id}/export?report_output_format=xml",
            headers=admin_auth_headers,
        )

        assert response.status_code == 400
        assert "no es soportado" in response.json().get("error", "").lower()

    def test_export_report_forbidden_for_player(self, client, player_auth_headers, usage_data):
        response = client.get("/api/v1/reports/game-usage/1/export", headers=player_auth_headers)
        assert response.status_code == 403

    @pytest.mark.parametrize("analysis", ["ranks", "characters", "teams"])
    def test_general_reports_persist_export_and_have_independent_history(self, client, admin_auth_headers, usage_data, analysis):
        game_id = usage_data["lol"].videogame_id
        if analysis != "teams":
            individual = self._generate(client, admin_auth_headers, game_id, analysis)
            assert individual.status_code == 201
        response = self._generate(client, admin_auth_headers, None, analysis)
        assert response.status_code == 201
        report = response.json()
        assert report["videogame_id"] is None
        assert report["videogame_name"] == "Todos los videojuegos"
        assert report["has_previous_report"] is False
        if analysis == "teams":
            assert [(item["name"], item["count"]) for item in report["items"]] == [("League of Legends", 2)]
        else:
            assert report["total"] == 3
            assert all(item["name"].startswith("League of Legends / ") for item in report["items"])
        for format in ["csv", "pdf", "excel"]:
            exported = client.get(f"/api/v1/reports/game-usage/{report['report_id']}/export",
                                  params={"report_output_format": format}, headers=admin_auth_headers)
            assert exported.status_code == 200
            assert "general" in exported.headers["content-disposition"]
            assert exported.content
        latest = client.get("/api/v1/reports/game-usage/last", params={"analysis_type": analysis}, headers=admin_auth_headers)
        assert latest.status_code == 200
        assert latest.json()["report_id"] == report["report_id"]
        reports = client.get("/api/v1/reports/game-usage", params={"general_only": True}, headers=admin_auth_headers).json()["reports"]
        assert [row["report_id"] for row in reports] == [report["report_id"]]
        second = self._generate(client, admin_auth_headers, None, analysis).json()
        assert second["previous_report_id"] == report["report_id"]
        assert all(item["variation"] == 0 for item in second["items"])

    @pytest.mark.parametrize("deleted_entity", ["profile", "user", "game", "rank", "role_profile", "character", "character_priority"])
    def test_usage_aggregates_exclude_soft_deleted_entities(self, client, admin_auth_headers, usage_data, test_db_session, deleted_entity):
        profile = usage_data["profiles"][0]
        entities = {
            "profile": profile,
            "user": usage_data["players"][0],
            "game": usage_data["lol"],
            "rank": usage_data["gold"],
            "role_profile": test_db_session.query(RoleProfileORM).filter_by(game_profile_id=profile.game_profile_id).first(),
            "character": usage_data["ahri"],
            "character_priority": test_db_session.query(CharacterPriorityORM).filter_by(game_profile_id=profile.game_profile_id).first(),
        }
        entities[deleted_entity].deleted = True
        test_db_session.commit()
        analysis = "characters" if deleted_entity in ("character", "character_priority") else "ranks"
        result = self._generate(client, admin_auth_headers, None, analysis)
        if deleted_entity == "game":
            assert result.status_code == 404
            return
        assert result.status_code == 201
        report = result.json()
        assert report["total"] == (2 if deleted_entity in ("profile", "user") else 3)
        counts = {item["name"].split(" / ")[-1]: item["count"] for item in report["items"]}
        if deleted_entity in ("rank", "character"):
            assert ("Gold" if analysis == "ranks" else "Ahri") not in counts
        else:
            assert counts["Gold" if analysis == "ranks" else "Ahri"] == 1

    def test_general_report_keeps_same_named_ranks_of_different_games_separate(self, client, admin_auth_headers, usage_data, test_db_session):
        game = usage_data["valorant"]
        rank = RankORM(name="Gold", value=100, videogame_id=game.videogame_id, icon_url="/gold.png")
        role = RoleORM(name="Duelist", videogame_id=game.videogame_id, icon_url="/role.png")
        profile = GameProfileORM(player_id=usage_data["players"][0].user_id, videogame_id=game.videogame_id)
        test_db_session.add_all([rank, role, profile])
        test_db_session.flush()
        test_db_session.add(RoleProfileORM(game_profile_id=profile.game_profile_id, role_id=role.role_id, rank_id=rank.rank_id))
        test_db_session.commit()
        response = client.post("/api/v1/reports/game-usage", json={"analysis_type": "ranks"}, headers=admin_auth_headers)
        assert response.status_code == 201
        report = response.json()
        assert report["total"] == 4
        assert {item["name"]: item["count"] for item in report["items"]} == {
            "League of Legends / Gold": 2, "League of Legends / Platinum": 1, "Valorant / Gold": 1,
        }

    def test_general_report_requires_admin(self, client, player_auth_headers):
        response = client.post("/api/v1/reports/game-usage", json={"analysis_type": "ranks"}, headers=player_auth_headers)
        assert response.status_code == 403
