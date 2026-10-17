import io
from datetime import datetime
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook

from app.application.use_cases.export_game_profile_report import ExportGameProfileReport
from app.domain.exceptions.report.report_not_found_exception import ReportNotFoundException
from app.domain.exceptions.report.unsupported_report_format_exception import UnsupportedReportFormatException
from app.domain.models.game_profile_report import GameProfileReport
from app.domain.models.game_profile_report_item import GameProfileReportItem


def _sample_report() -> GameProfileReport:
    items = [
        GameProfileReportItem(
            item_id=1, videogame_id=1, videogame_name="League of Legends", icon_url=None,
            profile_count=3, percentage=75.0, variation=1, variation_percentage=50.0, trend="sube",
        ),
        GameProfileReportItem(
            item_id=2, videogame_id=2, videogame_name="Valorant", icon_url=None,
            profile_count=1, percentage=25.0, variation=-1, variation_percentage=-50.0, trend="baja",
        ),
    ]
    return GameProfileReport(
        report_id=5,
        generated_by=1,
        generated_at=datetime(2026, 6, 30, 12, 0, 0),
        total_profiles=4,
        items=items,
        start_date=datetime(2026, 6, 1, 0, 0, 0),
        end_date=datetime(2026, 6, 30, 23, 59, 59),
        previous_report_id=4,
    )


class TestExportGameProfileReportUseCase:

    @pytest.fixture
    def mock_deps(self):
        uow = MagicMock()
        uow.__enter__.return_value = uow
        uow.__exit__.return_value = None
        uow.game_profile_report_repo = MagicMock()
        use_case = ExportGameProfileReport(unit_of_work=uow)
        return use_case, uow

    def test_export_csv(self, mock_deps):
        use_case, uow = mock_deps
        uow.game_profile_report_repo.get_report_by_id.return_value = _sample_report()

        exported = use_case.execute(5, "csv")

        assert exported.filename == "reporte_perfiles_por_videojuego_5.csv"
        assert exported.media_type == "text/csv"
        content = exported.content.decode("utf-8")
        assert "VIDEOJUEGO,PERFILES,PORCENTAJE,VARIACION,VARIACION_PORCENTAJE,TENDENCIA" in content
        assert "League of Legends,3,75.0,1,50.0,sube" in content
        assert "TOTAL,4,100.0" in content

    def test_export_pdf(self, mock_deps):
        use_case, uow = mock_deps
        uow.game_profile_report_repo.get_report_by_id.return_value = _sample_report()

        exported = use_case.execute(5, "pdf")

        assert exported.filename.endswith(".pdf")
        assert exported.media_type == "application/pdf"
        assert exported.content.startswith(b"%PDF")

    def test_export_excel(self, mock_deps):
        use_case, uow = mock_deps
        uow.game_profile_report_repo.get_report_by_id.return_value = _sample_report()

        exported = use_case.execute(5, "excel")

        assert exported.filename.endswith(".xlsx")
        assert "spreadsheetml" in exported.media_type
        worksheet = load_workbook(io.BytesIO(exported.content)).active
        values = [cell.value for row in worksheet.iter_rows() for cell in row]
        assert "League of Legends" in values
        assert "Valorant" in values

    def test_export_unsupported_format(self, mock_deps):
        use_case, _ = mock_deps
        with pytest.raises(UnsupportedReportFormatException):
            use_case.execute(5, "xml")

    def test_export_report_not_found(self, mock_deps):
        use_case, uow = mock_deps
        uow.game_profile_report_repo.get_report_by_id.return_value = None
        with pytest.raises(ReportNotFoundException):
            use_case.execute(999, "csv")
