import io

import pytest
from openpyxl import load_workbook

from app.application.services.report_export_service import (
    CSV_FORMAT,
    CSV_MEDIA_TYPE,
    EXCEL_FORMAT,
    EXCEL_MEDIA_TYPE,
    PDF_FORMAT,
    PDF_MEDIA_TYPE,
    ExportedReportFile,
    ReportExportService,
    ReportTable,
)
from app.domain.exceptions.report.unsupported_report_format_exception import UnsupportedReportFormatException


def _sample_table() -> ReportTable:
    return ReportTable(
        title="Reporte de prueba",
        meta_lines=["Videojuego: League of Legends", "Período: histórico acumulado"],
        headers=["CATEGORIA", "CANTIDAD", "PORCENTAJE", "VARIACION", "VARIACION %", "TENDENCIA"],
        rows=[
            ["Gold", 2, 66.67, 1, 100.0, "sube"],
            ["Platinum", 1, 33.33, None, None, None],
        ],
        total_row=["TOTAL", 3, 100.0, None, None, None],
        column_widths=[20, 12, 12, 12, 14, 12],
        pdf_decimal_columns=(2, 4),
    )


class TestReportExportService:

    # ---------------------------------------------------------
    # ensure_supported_format
    # ---------------------------------------------------------

    def test_ensure_supported_format_normalizes_case(self):
        assert ReportExportService.ensure_supported_format("PDF") == PDF_FORMAT
        assert ReportExportService.ensure_supported_format("Excel") == EXCEL_FORMAT

    def test_ensure_supported_format_defaults_to_csv_when_empty(self):
        assert ReportExportService.ensure_supported_format("") == CSV_FORMAT
        assert ReportExportService.ensure_supported_format(None) == CSV_FORMAT

    def test_ensure_supported_format_rejects_unknown(self):
        with pytest.raises(UnsupportedReportFormatException) as exc_info:
            ReportExportService.ensure_supported_format("xml")
        assert exc_info.value.status_code == 400

    # ---------------------------------------------------------
    # build
    # ---------------------------------------------------------

    def test_build_csv_returns_exported_file(self):
        exported = ReportExportService.build("csv", "mi_reporte", _sample_table())

        assert isinstance(exported, ExportedReportFile)
        assert exported.filename == "mi_reporte.csv"
        assert exported.media_type == CSV_MEDIA_TYPE
        content = exported.content.decode("utf-8")
        assert "CATEGORIA,CANTIDAD,PORCENTAJE,VARIACION,VARIACION %,TENDENCIA" in content
        assert "Gold,2,66.67,1,100.0,sube" in content
        assert "Platinum,1,33.33,,," in content
        assert "TOTAL,3,100.0,,," in content

    def test_build_pdf_starts_with_pdf_header(self):
        exported = ReportExportService.build("pdf", "mi_reporte", _sample_table())

        assert exported.filename == "mi_reporte.pdf"
        assert exported.media_type == PDF_MEDIA_TYPE
        assert exported.content.startswith(b"%PDF")

    def test_build_excel_is_a_workbook_with_rows(self):
        exported = ReportExportService.build("excel", "mi_reporte", _sample_table())

        assert exported.filename == "mi_reporte.xlsx"
        assert exported.media_type == EXCEL_MEDIA_TYPE
        assert exported.content.startswith(b"PK")
        worksheet = load_workbook(io.BytesIO(exported.content)).active
        values = [cell.value for row in worksheet.iter_rows() for cell in row]
        assert "Reporte de prueba" in values
        assert "Gold" in values
        assert "Platinum" in values
        assert "TOTAL" in values

    def test_build_rejects_unknown_format(self):
        with pytest.raises(UnsupportedReportFormatException):
            ReportExportService.build("xml", "mi_reporte", _sample_table())

    # ---------------------------------------------------------
    # helpers de formato de celdas
    # ---------------------------------------------------------

    def test_csv_cells_replaces_none_with_empty_string(self):
        assert ReportExportService._csv_cells(["A", None, 3]) == ["A", "", 3]

    def test_pdf_cells_formats_decimal_columns_and_empty_values(self):
        cells = ReportExportService._pdf_cells(["Gold", 2, 66.666, 1, 100.0, None], (2, 4))
        assert cells == ["Gold", "2", "66.67", "1", "100.00", ""]

    @pytest.mark.parametrize("name", ["=SUM(1,2)", "+1+1", "-1+1", "@SUM(A1)", "\t=1+1"])
    def test_csv_neutralizes_formula_strings_but_preserves_negative_numbers(self, name):
        assert ReportExportService._csv_cells([name, -3, None]) == ["'" + name, -3, ""]

    def test_excel_keeps_formula_like_catalog_text_literal_and_supports_wide_tables(self):
        table = _sample_table()
        table.rows[0][0] = "=1+1"
        table.headers = [f"Column {index}" for index in range(28)]
        table.column_widths = [15] * 28
        exported = ReportExportService.build("excel", "wide", table)
        sheet = load_workbook(io.BytesIO(exported.content), data_only=False).active
        header = next(cell.row for row in sheet for cell in row if cell.value == "Column 0")
        assert sheet.cell(header + 1, 1).value == "=1+1"
        assert sheet.cell(header + 1, 1).data_type == "s"
        assert sheet.cell(header + 1, 4).value == 1
        assert sheet.cell(header + 1, 3).number_format == "0.00"
        assert sheet.freeze_panes == f"B{header + 1}"
        assert sheet.auto_filter.ref == f"A{header}:AB{header + 2}"
        assert sheet.column_dimensions["AB"].width == 15

    def test_pdf_handles_markup_in_catalog_names_and_long_labels(self):
        table = _sample_table()
        table.title = "Report <script> & names"
        table.meta_lines.append("Game: <b>untrusted & text</b>")
        table.rows[0][0] = "Game <img src='file:///etc/passwd'> & " + "long name " * 15
        table.headers[0] = "Category <test>"
        assert ReportExportService.build("pdf", "escaped", table).content.startswith(b"%PDF")
