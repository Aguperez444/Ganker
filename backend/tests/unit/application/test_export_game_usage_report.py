import io
from datetime import datetime
from unittest.mock import MagicMock

import pytest
from openpyxl import load_workbook

from app.application.use_cases.export_game_usage_report import ExportGameUsageReport
from app.domain.exceptions.report.report_not_found_exception import ReportNotFoundException
from app.domain.exceptions.report.unsupported_report_format_exception import UnsupportedReportFormatException
from app.domain.models.game_usage_report import GameUsageReport
from app.domain.models.game_usage_report_item import GameUsageReportItem


def _sample_report(analysis_type: str = "ranks") -> GameUsageReport:
    items = [
        GameUsageReportItem(
            item_id=1, category_id=1, name="Gold", icon_url=None,
            count=2, percentage=66.67, variation=1, variation_percentage=100.0, trend="sube",
        ),
        GameUsageReportItem(
            item_id=2, category_id=2, name="Platinum", icon_url=None,
            count=1, percentage=33.33, variation=0, variation_percentage=0.0, trend="estable",
        ),
    ]
    return GameUsageReport(
        report_id=7,
        generated_by=1,
        analysis_type=analysis_type,
        videogame_id=1,
        videogame_name="League of Legends",
        icon_url=None,
        generated_at=datetime(2026, 6, 30, 12, 0, 0),
        total=3,
        items=items,
        start_date=datetime(2026, 6, 1, 0, 0, 0),
        end_date=datetime(2026, 6, 30, 23, 59, 59),
        previous_report_id=6,
    )


class TestExportGameUsageReportUseCase:

    @pytest.fixture
    def mock_deps(self):
        uow = MagicMock()
        uow.__enter__.return_value = uow
        uow.__exit__.return_value = None
        uow.game_usage_report_repo = MagicMock()
        use_case = ExportGameUsageReport(unit_of_work=uow)
        return use_case, uow

    def test_export_csv(self, mock_deps):
        use_case, uow = mock_deps
        uow.game_usage_report_repo.get_report_by_id.return_value = _sample_report("ranks")

        exported = use_case.execute(7, "csv")

        assert exported.filename == "reporte_uso_ranks_videojuego_1.csv"
        assert exported.media_type == "text/csv"
        content = exported.content.decode("utf-8")
        assert "RANGO,JUGADORES,PORCENTAJE,VARIACION,VARIACION %,TENDENCIA" in content
        assert "Gold,2,66.67,1,100.0,sube" in content
        assert "Platinum,1,33.33,0,0.0,estable" in content
        assert "TOTAL,3,100.0" in content

    def test_export_characters_uses_its_headers(self, mock_deps):
        use_case, uow = mock_deps
        uow.game_usage_report_repo.get_report_by_id.return_value = _sample_report("characters")

        exported = use_case.execute(7, "csv")

        assert exported.filename == "reporte_uso_characters_videojuego_1.csv"
        content = exported.content.decode("utf-8")
        assert "PERSONAJE,VECES_ELEGIDO,PORCENTAJE" in content

    def test_export_teams_uses_its_headers(self, mock_deps):
        use_case, uow = mock_deps
        uow.game_usage_report_repo.get_report_by_id.return_value = _sample_report("teams")

        exported = use_case.execute(7, "csv")

        assert exported.filename == "reporte_uso_teams_videojuego_1.csv"
        content = exported.content.decode("utf-8")
        assert "FRANJA,EQUIPOS,PORCENTAJE" in content

    def test_export_pdf(self, mock_deps):
        use_case, uow = mock_deps
        uow.game_usage_report_repo.get_report_by_id.return_value = _sample_report("ranks")

        exported = use_case.execute(7, "pdf")

        assert exported.filename.endswith(".pdf")
        assert exported.media_type == "application/pdf"
        assert exported.content.startswith(b"%PDF")

    def test_export_excel(self, mock_deps):
        use_case, uow = mock_deps
        uow.game_usage_report_repo.get_report_by_id.return_value = _sample_report("ranks")

        exported = use_case.execute(7, "excel")

        assert exported.filename.endswith(".xlsx")
        assert "spreadsheetml" in exported.media_type
        worksheet = load_workbook(io.BytesIO(exported.content)).active
        values = [cell.value for row in worksheet.iter_rows() for cell in row]
        assert "Gold" in values
        assert "Platinum" in values
        assert "TOTAL" in values

    def test_export_unsupported_format(self, mock_deps):
        use_case, _ = mock_deps

        with pytest.raises(UnsupportedReportFormatException):
            use_case.execute(7, "xml")

    def test_export_report_not_found(self, mock_deps):
        use_case, uow = mock_deps
        uow.game_usage_report_repo.get_report_by_id.return_value = None

        with pytest.raises(ReportNotFoundException):
            use_case.execute(999, "csv")
