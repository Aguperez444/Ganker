import csv
import io
from dataclasses import dataclass
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.domain.exceptions.report.unsupported_report_format_exception import UnsupportedReportFormatException

CSV_FORMAT = "csv"
PDF_FORMAT = "pdf"
EXCEL_FORMAT = "excel"

CSV_MEDIA_TYPE = "text/csv"
PDF_MEDIA_TYPE = "application/pdf"
EXCEL_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

SUPPORTED_FORMATS = (CSV_FORMAT, PDF_FORMAT, EXCEL_FORMAT)

_REPORT_AUTHOR = "Ganker"
_EXCEL_SHEET_TITLE = "Reporte"

# Estilo visual común de la tabla en PDF (mismo formato para todos los reportes).
_PDF_TABLE_STYLE = (
    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2C3E50")),
    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
    ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
    ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#ECEFF1")),
    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
    ("ALIGN", (1, 1), (-1, -1), "CENTER"),
    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ("FONTSIZE", (0, 0), (-1, -1), 9),
)


@dataclass
class ExportedReportFile:
    """Archivo exportado listo para descargar, independiente del tipo de reporte."""

    filename: str
    content: bytes
    media_type: str


@dataclass
class ReportTable:
    """
    Representación tabular neutra de un reporte, lista para exportar. El dominio de cada reporte se
    traduce a esta estructura y el servicio se encarga del formato de salida (CSV, PDF o Excel).
    """

    title: str
    meta_lines: list[str]
    headers: list[str]
    rows: list[list]
    total_row: list
    column_widths: list[int]
    pdf_decimal_columns: tuple[int, ...] = ()


class ReportExportService:
    """Centraliza el render de exportación (CSV, PDF y Excel) común a todos los reportes."""

    @staticmethod
    def ensure_supported_format(report_format: str) -> str:
        normalized = (report_format or CSV_FORMAT).strip().lower()
        if normalized not in SUPPORTED_FORMATS:
            raise UnsupportedReportFormatException(normalized)
        return normalized

    @classmethod
    def build(cls, report_format: str, base_filename: str, table: ReportTable) -> ExportedReportFile:
        report_format = cls.ensure_supported_format(report_format)

        if report_format == CSV_FORMAT:
            return ExportedReportFile(f"{base_filename}.csv", cls._build_csv(table), CSV_MEDIA_TYPE)
        if report_format == PDF_FORMAT:
            return ExportedReportFile(f"{base_filename}.pdf", cls._build_pdf(table), PDF_MEDIA_TYPE)
        return ExportedReportFile(f"{base_filename}.xlsx", cls._build_excel(table), EXCEL_MEDIA_TYPE)

    # ---------------------------------------------------------
    # CSV
    # ---------------------------------------------------------

    @classmethod
    def _build_csv(cls, table: ReportTable) -> bytes:
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(cls._csv_cells(table.headers))
        for row in table.rows:
            writer.writerow(cls._csv_cells(row))
        writer.writerow([])
        writer.writerow(cls._csv_cells(table.total_row))
        return buffer.getvalue().encode("utf-8")

    @staticmethod
    def _csv_cells(row: list) -> list:
        # Spreadsheet programs interpret these string prefixes as formulas.
        return [
            "" if cell is None else
            "'" + cell if isinstance(cell, str) and cell.lstrip().startswith(("=", "+", "-", "@")) else cell
            for cell in row
        ]

    # ---------------------------------------------------------
    # PDF
    # ---------------------------------------------------------

    @classmethod
    def _build_pdf(cls, table: ReportTable) -> bytes:
        buffer = io.BytesIO()
        page_size = landscape(A4) if len(table.headers) > 4 else A4
        document = SimpleDocTemplate(buffer, pagesize=page_size, title=table.title, author=_REPORT_AUTHOR,
                                     leftMargin=1.2 * cm, rightMargin=1.2 * cm,
                                     topMargin=1.3 * cm, bottomMargin=1.5 * cm)
        styles = getSampleStyleSheet()

        elements = [Paragraph(escape(table.title), styles["Title"])]
        elements += [Paragraph(escape(line), styles["Normal"]) for line in table.meta_lines]
        elements.append(Spacer(1, 0.5 * cm))

        data = [table.headers]
        data += [cls._pdf_cells(row, table.pdf_decimal_columns) for row in table.rows]
        data.append(cls._pdf_cells(table.total_row, table.pdf_decimal_columns))

        # Paragraphs wrap long names and escape catalog text interpreted as markup.
        cell_style = ParagraphStyle("ReportCell", parent=styles["Normal"], fontSize=8, leading=10)
        header_style = ParagraphStyle("ReportHeader", parent=cell_style, textColor=colors.white,
                                      fontName="Helvetica-Bold")
        data = [[Paragraph(escape(str(value)), header_style if row_index == 0 else cell_style)
                 for value in row] for row_index, row in enumerate(data)]
        weights = table.column_widths or [1] * len(table.headers)
        col_widths = [document.width * width / sum(weights) for width in weights]
        pdf_table = Table(data, colWidths=col_widths, repeatRows=1, hAlign="LEFT")
        pdf_table.setStyle(TableStyle(_PDF_TABLE_STYLE))
        elements.append(pdf_table)

        def footer(canvas, doc):
            canvas.saveState()
            canvas.setFont("Helvetica", 8)
            canvas.setFillColor(colors.HexColor("#526174"))
            canvas.drawString(doc.leftMargin, 0.7 * cm, "Ganker | Administración")
            canvas.drawRightString(page_size[0] - doc.rightMargin, 0.7 * cm, f"Página {doc.page}")
            canvas.restoreState()

        document.build(elements, onFirstPage=footer, onLaterPages=footer)
        return buffer.getvalue()

    @staticmethod
    def _pdf_cells(row: list, decimal_columns: tuple[int, ...]) -> list[str]:
        cells = []
        for index, cell in enumerate(row):
            if cell is None or cell == "":
                cells.append("")
            elif index in decimal_columns and isinstance(cell, float):
                cells.append(f"{cell:.2f}")
            else:
                cells.append(str(cell))
        return cells

    # ---------------------------------------------------------
    # Excel
    # ---------------------------------------------------------

    @classmethod
    def _build_excel(cls, table: ReportTable) -> bytes:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = _EXCEL_SHEET_TITLE

        worksheet.append([table.title])
        for line in table.meta_lines:
            worksheet.append([line])
        worksheet.append([])

        worksheet.append(table.headers)
        header_row = worksheet.max_row
        for cell in worksheet[header_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="2C3E50")
            cell.alignment = Alignment(wrap_text=True)
        worksheet.row_dimensions[header_row].height = 30

        for row in table.rows:
            worksheet.append(row)
        worksheet.append(table.total_row)

        last_column = get_column_letter(len(table.headers))
        worksheet.freeze_panes = f"B{header_row + 1}"
        worksheet.auto_filter.ref = f"A{header_row}:{last_column}{worksheet.max_row - 1}"
        worksheet.print_title_rows = f"{header_row}:{header_row}"
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A4
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.oddFooter.center.text = "Ganker | Página &P de &N"
        worksheet["A1"].font = Font(size=18, bold=True, color="2C3E50")
        for row_index in range(1, header_row - 1):
            worksheet.merge_cells(start_row=row_index, start_column=1,
                                  end_row=row_index, end_column=len(table.headers))
            worksheet.cell(row_index, 1).alignment = Alignment(wrap_text=True, vertical="top")
            worksheet.row_dimensions[row_index].height = 32
        for row in worksheet.iter_rows():
            for cell in row:
                # Keep catalog text literal, including strings beginning with '='.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                if cell.row > header_row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
                    if cell.column - 1 in table.pdf_decimal_columns and isinstance(cell.value, (int, float)):
                        cell.number_format = '0.00'
                    if cell.row == worksheet.max_row:
                        cell.font = Font(bold=True)
                        cell.fill = PatternFill("solid", fgColor="ECEFF1")
        for column_index, width in enumerate(table.column_widths, start=1):
            worksheet.column_dimensions[get_column_letter(column_index)].width = width

        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()
